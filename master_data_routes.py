"""Buyers, Products and Business Profiles pages, and their Excel import.

Pages
    /buyers, /products, /business-profiles

JSON (read)
    /api/buyers/overview            buyers + per-buyer sales + summary tiles
    /api/products/overview          products + per-product sales + summary tiles
    /api/business-profiles/overview profiles + invoices issued per profile

Excel
    /api/buyers/import-template     sample .xlsx to fill in
    /api/products/import-template
    /api/buyers/import              POST .xlsx; commit=0 checks, commit=1 saves
    /api/products/import

Creating, editing and deleting single records still goes through the
endpoints in invoice_form_routes.py, which the invoice form also uses; the
rules they share (codes, duplicates) are in master_data.py.

Sales figures come from this client's *submitted* invoices in the current
environment, summed with Decimal. Everything is scoped to session client_id.
"""

import io
import json
import random
from datetime import date, datetime, timedelta
from decimal import Decimal

from flask import jsonify, render_template, request, send_file, session, current_app

import master_data as md
from fbr_reference_routes import canonical_province, fetch_provinces

MAX_IMPORT_ROWS = 2000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --------------------------------------------------------------------- helpers

def _as_dict(raw):
    if isinstance(raw, dict):
        return raw
    if isinstance(raw, str):
        try:
            value = json.loads(raw)
            return value if isinstance(value, dict) else {}
        except ValueError:
            return {}
    return {}


def _dec(value):
    return md.to_decimal(value, Decimal("0"))


def _money(value):
    return str(value.quantize(Decimal("0.01")))


def _plain(value):
    """Decimal without trailing zeros or exponent: 18.00 -> "18", 100 -> "100"."""
    text = format(value.normalize(), "f")
    return text


def _invoice_day(data, created_at):
    text = str(data.get("invoiceDate") or "")[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return created_at.date() if created_at else None


def _submitted_invoices(cur, client_id, env):
    """(invoice_data dict, invoice date) for each accepted invoice."""
    cur.execute(
        """
        SELECT invoice_data, created_at
        FROM invoices
        WHERE client_id = %s AND env = %s AND LOWER(COALESCE(status, '')) = 'success'
        """,
        (client_id, env),
    )
    for raw, created_at in cur.fetchall():
        data = _as_dict(raw)
        yield data, _invoice_day(data, created_at)


def _line_total(item):
    return _dec(item.get("valueSalesExcludingST")) + _dec(item.get("salesTaxApplicable"))


def _form_options():
    """The same option lists the invoice form's dropdowns use."""
    try:
        view = current_app.view_functions.get("get_form_options")
        if view:
            response = view()
            if isinstance(response, tuple):
                response = response[0]
            data = response.get_json(silent=True) or {}
            if isinstance(data, dict):
                return data
    except Exception as exc:  # FBR down etc. -- fall back to what we know
        print(f"[master-data] form options unavailable: {exc}")
    return {}


def _cell_text(value):
    """Excel cell -> clean string. Whole-number floats lose their '.0' so an
    NTN typed as a number (1234567.0) reads back as 1234567."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return " ".join(str(value).split())


def _header_key(text):
    return "".join(ch for ch in str(text or "").lower() if ch.isalnum())


def _read_sheet(upload, aliases):
    """Parse the first worksheet: header row -> field names via `aliases`."""
    import openpyxl  # imported lazily: only the import path needs it

    raw = upload.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise ValueError("The file is larger than 5 MB.")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError("That file could not be read. Upload the .xlsx template, filled in.")
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ValueError("The sheet is empty.")
    columns = {}
    for index, title in enumerate(header):
        field = aliases.get(_header_key(title))
        if field and field not in columns:
            columns[field] = index
    out = []
    for number, row in enumerate(rows, start=2):
        values = {field: _cell_text(row[i] if i < len(row) else None) for field, i in columns.items()}
        if not any(values.values()):
            continue
        out.append((number, values))
        if len(out) > MAX_IMPORT_ROWS:
            raise ValueError(f"Import at most {MAX_IMPORT_ROWS} rows at a time.")
    return columns, out


def _style_sheet(ws, headers, widths, required):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.append(headers)
    fill = PatternFill("solid", fgColor="273FA8")
    optional_fill = PatternFill("solid", fgColor="5C6480")
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill if title in required else optional_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = widths[col - 1]
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _put_row(ws, values):
    """Write the next sample row under the header. Not ws.append(): the text
    format pre-applied to 2,000 rows makes append() start after them."""
    row = getattr(ws, "_tlp_next_row", 2)
    for col, value in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=value)
    ws._tlp_next_row = row + 1


def _add_list_validation(wb, ws, column_letter, values, list_col):
    """Dropdown in `column_letter` (rows 2-2001) fed from a hidden Lists sheet."""
    from openpyxl.worksheet.datavalidation import DataValidation

    if not values:
        return
    lists = wb["Lists"] if "Lists" in wb.sheetnames else wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    for i, value in enumerate(values, start=1):
        lists[f"{list_col}{i}"] = value
    rule = DataValidation(type="list", formula1=f"=Lists!${list_col}$1:${list_col}${len(values)}",
                          allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(rule)
    rule.add(f"{column_letter}2:{column_letter}{MAX_IMPORT_ROWS + 1}")


def _add_instructions(wb, lines):
    from openpyxl.styles import Font

    info = wb.create_sheet("Instructions")
    info.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        info[f"A{i}"] = line
        if i == 1:
            info[f"A{i}"].font = Font(bold=True, size=13)


def _xlsx_response(wb, filename):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, mimetype=XLSX_MIME, as_attachment=True, download_name=filename)


# Sample values for the templates -- random each download, and plainly
# examples (the instructions say to replace them).
_SAMPLE_BUYERS = [
    "Al-Noor Traders", "Crescent Distributors", "Indus Medical Store", "Karachi Pharma House",
    "Lahore Surgical Co.", "Margalla Enterprises", "Paramount Stores", "Ravi Textiles (Pvt) Ltd",
    "Shifa Medical Distributors", "Sunrise Trading Company", "Unity Hardware Mart", "Zam Zam Foods",
]
_SAMPLE_STREETS = [
    "Plot 12, SITE Area, Karachi", "45-B Gulberg III, Lahore", "Office 7, Blue Area, Islamabad",
    "Shop 3, Saddar, Rawalpindi", "22 Mall Road, Peshawar", "House 9, Satellite Town, Quetta",
    "Unit 4, Industrial Estate, Faisalabad", "18 Civil Lines, Multan",
]
_SAMPLE_PRODUCTS = [
    ("Paracetamol 500mg Tablets (10x10)", "3004.9099", "BOX", "320.00"),
    ("Portland Cement 50kg", "2523.2900", "BAG", "1450.00"),
    ("Cotton Yarn 20s", "5205.1200", "KG", "780.50"),
    ("LED Bulb 12W", "8539.5200", "Numbers, pieces, units", "425.00"),
    ("Cooking Oil 5L", "1512.1900", "CAN", "2650.00"),
    ("Steel Bar 12mm", "7214.2000", "MT", "245000.00"),
    ("Mineral Water 1.5L (Carton of 6)", "2201.1010", "CARTON", "540.00"),
    ("PVC Pipe 2 inch", "3917.2300", "MTR", "310.00"),
]


def add_master_data_routes(app, get_db_connection, get_env):

    def _env():
        return session.get("env", "sandbox")

    def _username(cur, client_id):
        cur.execute(
            "SELECT u.username FROM users u JOIN clients c ON u.id = c.user_id WHERE c.id = %s",
            (client_id,),
        )
        row = cur.fetchone()
        return (row[0] or "").strip() if row else ""

    def _is_special(username):
        from invoice_form_routes import _is_special_username
        return _is_special_username(username)

    def _province_lookup():
        provinces, _ = fetch_provinces(get_db_connection, get_env)
        known = {(p.get("stateProvinceDesc") or "").strip() for p in provinces}

        def lookup(value):
            result = canonical_province(value, provinces)
            return result if result in known else ""
        return lookup, sorted(known)

    # ------------------------------------------------------------------ pages

    @app.route("/buyers")
    def buyers_page():
        return render_template("buyers.html", active_page="buyers",
                               query=(request.args.get("q") or "").strip()[:100])

    @app.route("/products")
    def products_page():
        return render_template("products.html", active_page="products",
                               query=(request.args.get("q") or "").strip()[:100])

    @app.route("/business-profiles")
    def business_profiles_page():
        return render_template("business-profiles.html", active_page="business-profiles")

    # ---------------------------------------------------------------- buyers

    @app.route("/api/buyers/overview", methods=["GET"])
    def buyers_overview():
        client_id = session.get("client_id")
        if not client_id:
            return jsonify({"error": "No client ID in session"}), 401
        conn = get_db_connection()
        cur = conn.cursor()
        try:
            cur.execute(
                """
                SELECT id, business_name, address, province, ntn_cnic, strn,
                       registration_type, buyer_code, is_default
                FROM buyers WHERE client_id = %s
                ORDER BY business_name
                """,
                (client_id,),
            )
            buyers = [
                {"id": r[0], "business_name": r[1] or "", "address": r[2] or "", "province": r[3] or "",
                 "ntn_cnic": r[4] or "", "strn": r[5] or "", "registration_type": r[6] or "",
                 "buyer_code": r[7] or "", "is_default": bool(r[8]),
                 "invoices": 0, "revenue": Decimal("0"), "last_invoice": None}
                for r in cur.fetchall()
            ]
            by_ntn = {md.normalize_tax_id(b["ntn_cnic"]): b for b in buyers if b["ntn_cnic"]}
            by_name = {md.name_key(b["business_name"]): b for b in buyers}

            today = date.today()
            total_revenue = Decimal("0")
            active = set()
            for data, day in _submitted_invoices(cur, client_id, _env()):
                total = sum((_line_total(i) for i in data.get("items") or [] if isinstance(i, dict)), Decimal("0"))
                total_revenue += total
                buyer = by_ntn.get(md.normalize_tax_id(data.get("buyerNTNCNIC"))) \
                    or by_name.get(md.name_key(data.get("buyerBusinessName")))
                if not buyer:
                    continue
                buyer["invoices"] += 1
                buyer["revenue"] += total
                if day and (buyer["last_invoice"] is None or day > buyer["last_invoice"]):
                    buyer["last_invoice"] = day
                if day and day.year == today.year and day.month == today.month:
                    active.add(buyer["id"])

            for b in buyers:
                b["revenue"] = _money(b["revenue"])
                b["last_invoice"] = b["last_invoice"].isoformat() if b["last_invoice"] else ""

            return jsonify({
                "buyers": buyers,
                "next_code": md.peek_next_code(cur, "buyers", "buyer_code", md.BUYER_CODE_PREFIX, client_id),
                "summary": {
                    "total": len(buyers),
                    "registered": sum(1 for b in buyers if b["registration_type"].lower() == "registered"),
                    "revenue": _money(total_revenue),
                    "active_this_month": len(active),
                },
            })
        finally:
            cur.close()
            conn.close()

    BUYER_ALIASES = {
        "businessname": "business_name", "buyername": "business_name", "name": "business_name",
        "buyer": "business_name",
        "ntncnic": "ntn_cnic", "ntn": "ntn_cnic", "cnic": "ntn_cnic", "ntnorcnic": "ntn_cnic",
        "strn": "strn", "salestaxregistrationno": "strn",
        "province": "province",
        "registrationtype": "registration_type", "type": "registration_type", "registration": "registration_type",
        "address": "address",
    }

    @app.route("/api/buyers/import-template", methods=["GET"])
    def buyers_import_template():
        import openpyxl

        _, provinces = _province_lookup()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Buyers"
        headers = ["Business Name*", "NTN/CNIC*", "STRN", "Province*", "Registration Type*", "Address*"]
        _style_sheet(ws, headers, [34, 18, 22, 24, 20, 46],
                     {h for h in headers if h.endswith("*")})
        for col in ("B", "C"):  # keep NTNs and STRNs as text so Excel doesn't mangle them
            for row in range(2, MAX_IMPORT_ROWS + 2):
                ws[f"{col}{row}"].number_format = "@"
        for name in random.sample(_SAMPLE_BUYERS, 5):
            registered = random.random() < 0.7
            ntn = str(random.randint(1000000, 9999999)) if registered or random.random() < 0.5 \
                else f"{random.randint(1000000000000, 9999999999999)}"
            strn = f"{random.randint(10, 99)}-{random.randint(10, 99)}-{random.randint(1000, 9999)}-{random.randint(100, 999)}-{random.randint(10, 99)}" if registered else ""
            _put_row(ws, [name, ntn, strn, random.choice(provinces) if provinces else "SINDH",
                          "Registered" if registered else "Unregistered", random.choice(_SAMPLE_STREETS)])
        _add_list_validation(wb, ws, "D", provinces, "A")
        _add_list_validation(wb, ws, "E", list(md.REGISTRATION_TYPES), "B")
        _add_instructions(wb, [
            "How to import buyers",
            "",
            "1. Replace the sample rows on the Buyers sheet with your own buyers (the samples are random examples).",
            "2. Columns marked * are required. STRN is optional.",
            "3. NTN/CNIC: a 7-character NTN or a 13-digit CNIC (dashes are fine).",
            "4. Province and Registration Type: pick from the dropdown in each cell.",
            "5. Buyer codes (B-0001, B-0002, ...) are assigned automatically -- don't add a code column.",
            "6. A row whose NTN/CNIC or business name matches a buyer you already have is skipped, not duplicated.",
            "7. Save the file and upload it on the Buyers page. You'll see a check of every row before anything is saved.",
        ])
        wb.move_sheet("Instructions", offset=-(len(wb.sheetnames) - 2))
        return _xlsx_response(wb, "buyers-import-template.xlsx")

    @app.route("/api/buyers/import", methods=["POST"])
    def buyers_import():
        client_id = session.get("client_id")
        if not client_id:
            return jsonify({"error": "No client ID in session"}), 401
        upload = request.files.get("file")
        if not upload:
            return jsonify({"error": "Choose the filled-in template to upload."}), 400
        commit = request.form.get("commit") == "1"
        try:
            columns, rows = _read_sheet(upload, BUYER_ALIASES)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        missing = [label for field, label in (("business_name", "Business Name"), ("ntn_cnic", "NTN/CNIC"))
                   if field not in columns]
        if missing:
            return jsonify({"error": "This doesn't look like the buyers template: no "
                                     + " or ".join(missing) + " column."}), 400

        lookup, _ = _province_lookup()
        conn = get_db_connection()
        cur = conn.cursor()
        try:
            existing = md.load_buyer_keys(cur, client_id)
            report, ready = [], []
            for number, values in rows:
                clean, error = md.clean_buyer(values, lookup)
                name = values.get("business_name") or ""
                if error:
                    if "Province" in error and values.get("province"):
                        error = f"“{values['province']}” is not a province FBR recognises."
                    report.append({"row": number, "name": name, "status": "error", "message": error})
                    continue
                clash = md.find_duplicate_buyer(existing, clean["business_name"], clean["ntn_cnic"])
                if clash:
                    report.append({"row": number, "name": name, "status": "duplicate",
                                   "message": md.buyer_duplicate_message(clash)})
                    continue
                # Later rows are checked against earlier rows of the same file.
                existing.append((None, clean["business_name"], clean["ntn_cnic"], f"row {number} of this file"))
                ready.append((number, clean))
                report.append({"row": number, "name": clean["business_name"], "status": "ok", "message": ""})

            added = 0
            if commit and ready:
                codes = md.reserve_codes(cur, "buyers", "buyer_code", md.BUYER_CODE_PREFIX, client_id, len(ready))
                by_row = {r["row"]: r for r in report}
                for (number, clean), code in zip(ready, codes):
                    cur.execute(
                        """
                        INSERT INTO buyers (client_id, business_name, address, province, ntn_cnic,
                                            strn, registration_type, buyer_code, is_default)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,FALSE)
                        """,
                        (client_id, clean["business_name"], clean["address"], clean["province"],
                         clean["ntn_cnic"], clean["strn"], clean["registration_type"], code),
                    )
                    by_row[number]["status"] = "added"
                    by_row[number]["message"] = code
                    added += 1
                conn.commit()
            return jsonify(_import_summary(report, commit, added))
        except Exception as exc:
            conn.rollback()
            return jsonify({"error": f"Import failed, nothing was saved: {exc}"}), 500
        finally:
            cur.close()
            conn.close()

    # --------------------------------------------------------------- products

    def _product_columns(cur):
        cur.execute(
            """
            SELECT column_name FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = 'products'
            """
        )
        return {r[0] for r in cur.fetchall()}

    @app.route("/api/products/overview", methods=["GET"])
    def products_overview():
        client_id = session.get("client_id")
        if not client_id:
            return jsonify({"error": "No client ID in session"}), 401
        conn = get_db_connection()
        cur = conn.cursor()
        try:
            cols = _product_columns(cur)
            has_code = "product_code" in cols
            has_sro_item = "sro_item_serial_no" in cols
            select = ["id", "description", "hs_code", "rate", "uom", "default_tax_rate",
                      "sro_schedule_no", "sale_type",
                      "product_code" if has_code else "''::text AS product_code",
                      "sro_item_serial_no" if has_sro_item else "''::text AS sro_item_serial_no"]
            cur.execute(
                f"""
                SELECT {', '.join(select)} FROM products
                WHERE client_id = %s AND (is_active = TRUE OR is_active IS NULL)
                ORDER BY description
                """,
                (client_id,),
            )
            products = []
            for r in cur.fetchall():
                products.append({
                    "id": r[0], "description": r[1] or "", "hs_code": r[2] or "",
                    "rate": _money(md.rate_key(r[3])), "uom": r[4] or "",
                    "default_tax_rate": _plain(md.to_decimal(r[5], Decimal("0"))),
                    "sro_schedule_no": r[6] or "", "sale_type": r[7] or "",
                    "product_code": r[8] or "", "sro_item_serial_no": r[9] or "",
                    "times_sold": 0, "quantity": Decimal("0"), "revenue": Decimal("0"), "last_sold": None,
                })

            # Match invoice lines to products by name; when several products
            # share a name (same item, different rates) the line goes to the
            # one whose rate is nearest the line's unit price.
            groups = {}
            for p in products:
                groups.setdefault(md.name_key(p["description"]), []).append(p)

            today = date.today()
            cutoff = today - timedelta(days=30)
            for data, day in _submitted_invoices(cur, client_id, _env()):
                for item in data.get("items") or []:
                    if not isinstance(item, dict):
                        continue
                    group = groups.get(md.name_key(item.get("productDescription")))
                    if not group:
                        continue
                    qty = _dec(item.get("quantity"))
                    if len(group) == 1:
                        product = group[0]
                    else:
                        unit = _dec(item.get("valueSalesExcludingST")) / qty if qty else Decimal("0")
                        product = min(group, key=lambda p: abs(Decimal(p["rate"]) - unit))
                    product["times_sold"] += 1
                    product["quantity"] += qty
                    product["revenue"] += _line_total(item)
                    if day and (product["last_sold"] is None or day > product["last_sold"]):
                        product["last_sold"] = day

            catalog_revenue = sum((p["revenue"] for p in products), Decimal("0"))
            active = sum(1 for p in products if p["last_sold"] and p["last_sold"] >= cutoff)
            rates = [md.to_decimal(p["default_tax_rate"], Decimal("0")) for p in products]
            avg_rate = (sum(rates, Decimal("0")) / len(rates)) if rates else Decimal("0")
            for p in products:
                p["revenue"] = _money(p["revenue"])
                p["quantity"] = _plain(p["quantity"])
                p["last_sold"] = p["last_sold"].isoformat() if p["last_sold"] else ""

            special = _is_special(_username(cur, client_id))
            return jsonify({
                "products": products,
                "codes_enabled": has_code,
                "code_editable": special,
                "is_special": special,
                "next_code": md.peek_next_code(cur, "products", "product_code", md.PRODUCT_CODE_PREFIX, client_id)
                if has_code else "",
                "summary": {
                    "total": len(products),
                    "active_30d": active,
                    "avg_tax_rate": str(avg_rate.quantize(Decimal("0.1"))),
                    "revenue": _money(catalog_revenue),
                },
            })
        finally:
            cur.close()
            conn.close()

    PRODUCT_ALIASES = {
        "productname": "description", "name": "description", "description": "description",
        "product": "description", "itemname": "description",
        "hscode": "hs_code", "hs": "hs_code",
        "uom": "uom", "unit": "uom", "unitofmeasure": "uom",
        "rate": "rate", "unitprice": "rate", "price": "rate", "rateunit": "rate",
        "taxrate": "default_tax_rate", "tax": "default_tax_rate", "defaulttaxrate": "default_tax_rate",
        "salestaxrate": "default_tax_rate",
        "saletype": "sale_type",
        "sroscheduleno": "sro_schedule_no", "sroschedule": "sro_schedule_no",
        "sroitemserialno": "sro_item_serial_no", "sroitemno": "sro_item_serial_no",
    }

    def _option_values(options, key):
        return [str(o.get("value") or "") for o in (options.get(key) or []) if o.get("value")]

    def _uom_lookup(options):
        choices = options.get("uoms") or []
        table = {}
        for o in choices:
            value = str(o.get("value") or "")
            table[md.name_key(value)] = value
            table[md.name_key(o.get("label"))] = value
            table[md.name_key(str(o.get("label") or "").split(" - ")[0])] = value

        def lookup(value):
            return table.get(md.name_key(value), "")
        return lookup

    def _tax_rates(options):
        rates = set()
        for value in _option_values(options, "taxRates"):
            number = md.to_decimal(value)
            if number is not None:
                rates.add(number.normalize())
        return rates

    @app.route("/api/products/import-template", methods=["GET"])
    def products_import_template():
        import openpyxl

        options = _form_options()
        uoms = _option_values(options, "uoms")
        sale_types = _option_values(options, "saleTypes")
        tax_labels = [str(o.get("label") or "").replace("%", "") for o in options.get("taxRates") or []]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        headers = ["Product Name*", "HS Code", "UoM*", "Rate", "Tax Rate %*", "Sale Type",
                   "SRO Schedule No", "SRO Item Serial No"]
        _style_sheet(ws, headers, [38, 14, 24, 12, 13, 36, 22, 20], {h for h in headers if h.endswith("*")})
        for row in range(2, MAX_IMPORT_ROWS + 2):
            ws[f"B{row}"].number_format = "@"  # 3004.9099 must not become 3004.9099000001
        default_sale = next((s for s in sale_types if "standard rate" in s.lower()), sale_types[0] if sale_types else "")
        for description, hs, uom, rate in random.sample(_SAMPLE_PRODUCTS, 5):
            if uoms and uom not in uoms:
                uom = uoms[0]
            _put_row(ws, [description, hs, uom, float(rate), random.choice([18, 18, 18, 17, 5, 0]),
                          default_sale, "", ""])
        _add_list_validation(wb, ws, "C", uoms, "A")
        _add_list_validation(wb, ws, "E", tax_labels, "B")
        _add_list_validation(wb, ws, "F", sale_types, "C")
        _add_instructions(wb, [
            "How to import products",
            "",
            "1. Replace the sample rows on the Products sheet with your own products (the samples are random examples).",
            "2. Columns marked * are required: Product Name, UoM and Tax Rate %.",
            "3. UoM, Tax Rate % and Sale Type: pick from the dropdown in each cell -- these are FBR's own values.",
            "4. Rate is the default unit price (optional). HS Code looks like 3004.9099.",
            "5. Product codes (P-0001, P-0002, ...) are assigned automatically -- don't add a code column.",
            "6. A row with the same Product Name AND Rate as a product you already have is skipped, not duplicated.",
            "   The same name at a different rate is added as a separate product.",
            "7. Save the file and upload it on the Products page. You'll see a check of every row before anything is saved.",
        ])
        wb.move_sheet("Instructions", offset=-(len(wb.sheetnames) - 2))
        return _xlsx_response(wb, "products-import-template.xlsx")

    @app.route("/api/products/import", methods=["POST"])
    def products_import():
        client_id = session.get("client_id")
        if not client_id:
            return jsonify({"error": "No client ID in session"}), 401
        upload = request.files.get("file")
        if not upload:
            return jsonify({"error": "Choose the filled-in template to upload."}), 400
        commit = request.form.get("commit") == "1"
        try:
            columns, rows = _read_sheet(upload, PRODUCT_ALIASES)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        if "description" not in columns:
            return jsonify({"error": "This doesn't look like the products template: no Product Name column."}), 400

        options = _form_options()
        uom_lookup = _uom_lookup(options)
        allowed_rates = _tax_rates(options)
        sale_types = {md.name_key(s): s for s in _option_values(options, "saleTypes")}

        conn = get_db_connection()
        cur = conn.cursor()
        try:
            cols = _product_columns(cur)
            has_code = "product_code" in cols
            has_sro_item = "sro_item_serial_no" in cols
            special = _is_special(_username(cur, client_id))
            existing = md.load_product_keys(cur, client_id)
            report, ready = [], []
            for number, values in rows:
                # A tax cell formatted as a percentage arrives as 0.18.
                tax = md.to_decimal(values.get("default_tax_rate"))
                if tax is not None and 0 < tax < 1 and (tax * 100).normalize() in allowed_rates:
                    values["default_tax_rate"] = str(tax * 100)
                clean, error = md.clean_product(values, uom_lookup)
                name = values.get("description") or ""
                if not error and allowed_rates and clean["default_tax_rate"].normalize() not in allowed_rates:
                    error = f"{_plain(clean['default_tax_rate'])}% is not one of FBR's tax rates."
                if not error and clean["sale_type"]:
                    match = sale_types.get(md.name_key(clean["sale_type"]))
                    if sale_types and not match:
                        error = f"“{clean['sale_type']}” is not an FBR sale type."
                    elif match:
                        clean["sale_type"] = match
                if error:
                    if error == "UoM is required." and values.get("uom"):
                        error = f"“{values['uom']}” is not an FBR unit of measure."
                    report.append({"row": number, "name": name, "status": "error", "message": error})
                    continue
                clash = md.find_duplicate_product(existing, clean["description"], clean["rate"])
                if clash:
                    report.append({"row": number, "name": name, "status": "duplicate",
                                   "message": md.product_duplicate_message(clash)})
                    continue
                deleted = md.find_duplicate_product(existing, clean["description"], clean["rate"], active=False)
                existing.append((None, clean["description"], clean["rate"], True))
                ready.append((number, clean, deleted[0] if deleted else None))
                report.append({"row": number, "name": clean["description"], "status": "ok",
                               "message": "Restores a deleted product" if deleted else ""})

            added = 0
            if commit and ready:
                if special:  # same as create_product for these clients
                    for _, clean, _ in ready:
                        clean["sale_type"] = clean["sro_schedule_no"] = clean["sro_item_serial_no"] = ""
                new_rows = [r for r in ready if r[2] is None]
                codes = iter(md.reserve_codes(cur, "products", "product_code", md.PRODUCT_CODE_PREFIX,
                                              client_id, len(new_rows))
                             if has_code and not special and new_rows else [])
                by_row = {r["row"]: r for r in report}
                for number, clean, restore_id in ready:
                    fields = {
                        "description": clean["description"], "hs_code": clean["hs_code"],
                        "rate": clean["rate"], "uom": clean["uom"],
                        "default_tax_rate": clean["default_tax_rate"],
                        "sro_schedule_no": clean["sro_schedule_no"], "sale_type": clean["sale_type"],
                    }
                    if has_sro_item:
                        fields["sro_item_serial_no"] = clean["sro_item_serial_no"]
                    if restore_id:
                        sets = ", ".join(f"{k} = %s" for k in fields)
                        cur.execute(f"UPDATE products SET {sets}, is_active = TRUE WHERE id = %s AND client_id = %s",
                                    list(fields.values()) + [restore_id, client_id])
                        by_row[number]["status"] = "restored"
                    else:
                        code = next(codes, "")
                        if has_code:
                            fields["product_code"] = code
                        names = ", ".join(fields)
                        marks = ", ".join(["%s"] * len(fields))
                        cur.execute(f"INSERT INTO products (client_id, {names}, is_active) VALUES (%s, {marks}, TRUE)",
                                    [client_id] + list(fields.values()))
                        by_row[number]["status"] = "added"
                        by_row[number]["message"] = code
                    added += 1
                conn.commit()
            return jsonify(_import_summary(report, commit, added))
        except Exception as exc:
            conn.rollback()
            return jsonify({"error": f"Import failed, nothing was saved: {exc}"}), 500
        finally:
            cur.close()
            conn.close()

    # ------------------------------------------------------ business profiles

    @app.route("/api/business-profiles/overview", methods=["GET"])
    def business_profiles_overview():
        client_id = session.get("client_id")
        if not client_id:
            return jsonify({"error": "No client ID in session"}), 401
        conn = get_db_connection()
        cur = conn.cursor()
        try:
            cur.execute(
                """
                SELECT id, business_name, address, province, ntn_cnic, strn, is_default
                FROM business_profiles WHERE client_id = %s
                ORDER BY is_default DESC, business_name
                """,
                (client_id,),
            )
            profiles = [
                {"id": r[0], "business_name": r[1] or "", "address": r[2] or "", "province": r[3] or "",
                 "ntn_cnic": r[4] or "", "strn": r[5] or "", "is_default": bool(r[6]), "invoices": 0}
                for r in cur.fetchall()
            ]
            by_ntn = {}
            for p in profiles:
                by_ntn.setdefault(md.normalize_tax_id(p["ntn_cnic"]), p)
            for data, _ in _submitted_invoices(cur, client_id, _env()):
                profile = by_ntn.get(md.normalize_tax_id(data.get("sellerNTNCNIC")))
                if profile:
                    profile["invoices"] += 1
            return jsonify({"profiles": profiles})
        finally:
            cur.close()
            conn.close()


def _import_summary(report, committed, added):
    counts = {"ok": 0, "duplicate": 0, "error": 0}
    for row in report:
        key = "ok" if row["status"] in ("ok", "added", "restored") else row["status"]
        counts[key] = counts.get(key, 0) + 1
    return {"rows": report, "committed": committed, "added": added,
            "summary": {"total": len(report), **counts}}
